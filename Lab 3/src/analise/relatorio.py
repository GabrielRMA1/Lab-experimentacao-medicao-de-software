"""Gera o relatório no modelo do professor (``docs/Template_Relatorio_Laboratorio.docx``).

Uso: ``python -m src.pipeline relatorio`` (depois de ``analise``).

O gerador abre o template, apaga os parágrafos de ORIENTAÇÃO e os marcadores
``[conteúdo do grupo ...]``, e escreve cada seção com os números reais de
``data/resultados`` (análise), ``data/funil.json`` (coleta) e
``data/validacao`` (amostra-ouro). O que só o grupo sabe (nomes, link,
Issues, print do Kanban, achados inesperados) sai **destacado em amarelo**
como ``[PREENCHER: ...]``. Revise o texto antes de entregar.
"""

from __future__ import annotations

import csv
import json
import re
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from src.analise.comum import ROOT, fmt, fmt_p
from src.metricas.dora import (
    NOMES, nivel_cfr, nivel_deploy_freq, nivel_lead_time, nivel_recuperacao,
)

TEMPLATE_PADRAO = ROOT / "docs" / "Template_Relatorio_Laboratorio.docx"
MARCADOR = re.compile(r"\[\[(.+?)\]\]")


# ------------------------------------------------------------------ docx

class Montador:
    """Insere parágrafos, listas, tabelas e figuras logo depois de uma âncora."""

    def __init__(self, doc: Document):
        self.doc = doc
        self.ancora = None
        self.n_figura = 0
        self.n_tabela = 0

    def em(self, elemento) -> "Montador":
        self.ancora = elemento
        return self

    def _encaixar(self, elemento) -> None:
        self.ancora.addnext(elemento)
        self.ancora = elemento

    @staticmethod
    def escrever(paragrafo, texto: str, negrito: bool = False) -> None:
        """Texto com ``[[...]]`` virando marcador amarelo ``[PREENCHER: ...]``."""

        pos = 0
        for m in MARCADOR.finditer(texto):
            if m.start() > pos:
                paragrafo.add_run(texto[pos:m.start()]).bold = negrito or None
            run = paragrafo.add_run(f"[PREENCHER: {m.group(1)}]")
            run.bold = True
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            pos = m.end()
        if pos < len(texto):
            paragrafo.add_run(texto[pos:]).bold = negrito or None

    def p(self, texto: str, estilo: str = "Normal", negrito: bool = False):
        par = self.doc.add_paragraph(style=estilo)
        self.escrever(par, texto, negrito)
        self._encaixar(par._p)
        return par

    def item(self, texto: str):
        return self.p(texto, "List Bullet")

    def h3(self, texto: str):
        return self.p(texto, "Heading 3")

    def legenda(self, texto: str):
        par = self.p(texto, "Caption")
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return par

    def tabela(self, cabecalho: list[str], linhas: list[list[str]], titulo: str,
               larguras_cm: list[float] | None = None) -> None:
        self.n_tabela += 1
        self.legenda(f"Tabela {self.n_tabela}. {titulo}")
        t = self.doc.add_table(rows=1 + len(linhas), cols=len(cabecalho))
        t.style = "Table Grid"
        for j, valor in enumerate(cabecalho):
            celula = t.rows[0].cells[j]
            celula.text = ""
            self.escrever(celula.paragraphs[0], str(valor), negrito=True)
            _sombrear(celula, "E8E8E8")
        for i, linha in enumerate(linhas, start=1):
            for j, valor in enumerate(linha):
                celula = t.rows[i].cells[j]
                celula.text = ""
                self.escrever(celula.paragraphs[0], "" if valor is None else str(valor))
        for linha in t.rows:
            for j, celula in enumerate(linha.cells):
                for par in celula.paragraphs:
                    par.paragraph_format.space_after = Pt(0)
                    for run in par.runs:
                        run.font.size = Pt(8.5)
                if larguras_cm:
                    celula.width = Cm(larguras_cm[j])
        if larguras_cm:  # Word lê a largura da célula; LibreOffice, a grade da tabela
            t.autofit = False
            for coluna, largura in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), larguras_cm):
                coluna.set(qn("w:w"), str(int(Cm(largura).twips)))
        self._encaixar(t._tbl)
        espaco = self.doc.add_paragraph()
        self._encaixar(espaco._p)

    def figura(self, caminho: Path, titulo: str, largura_cm: float = 16.5) -> None:
        self.n_figura += 1
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caminho.exists():
            par.add_run().add_picture(str(caminho), width=Cm(largura_cm))
        else:
            self.escrever(par, f"[[figura {caminho.name} não encontrada: rode `analise`]]")
        self._encaixar(par._p)
        self.legenda(f"Figura {self.n_figura}. {titulo}")


def _sombrear(celula, cor: str) -> None:
    tc_pr = celula._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    tc_pr.append(sombra)


def _texto(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _estilo(el) -> str:
    estilo = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return estilo.get(qn("w:val")) if estilo is not None else ""


def limpar_template(doc: Document) -> dict[str, Any]:
    """Remove orientações, marcadores e tabelas-exemplo; devolve os títulos.

    Mantém o título, a capa (primeira tabela) e os títulos numerados
    ("1. Introdução", "3.1 Principais Desafios"...). Tudo o mais, a partir de
    "1. Introdução", é apagado para receber o conteúdo do grupo.
    """

    corpo = doc.element.body
    titulos: dict[str, Any] = {}
    dentro = False
    for el in list(corpo.iterchildren()):
        if el.tag == qn("w:sectPr"):
            continue
        texto = _texto(el).strip()
        if el.tag == qn("w:p") and texto.startswith("1. Introdução"):
            dentro = True
        if not dentro:
            if texto.startswith("Este documento é um MODELO"):
                corpo.remove(el)
            continue
        estilo = _estilo(el).lower()
        if (el.tag == qn("w:p") and ("heading" in estilo or "ttulo" in estilo)
                and re.match(r"^\d+(\.\d+)?\.?\s", texto)):
            titulos[texto] = el
            continue
        corpo.remove(el)  # orientação, lista de perguntas, marcador, tabela-exemplo
    return titulos


def _titulo(titulos: dict[str, Any], prefixo: str):
    for texto, el in titulos.items():
        if texto.startswith(prefixo):
            return el
    raise KeyError(f"Título '{prefixo}' não encontrado no template.")


def preencher_capa(doc: Document, dados: dict[str, str]) -> None:
    for par in doc.paragraphs[:3]:
        if "Modelo/Template" in par.text:
            for run in par.runs[1:]:
                run.text = ""
            par.runs[0].text = "Laboratório de Experimentação de Software — Relatório do Lab03"
    tabela = doc.tables[0]
    for linha in tabela.rows:
        chave = linha.cells[0].text.strip()
        for prefixo, valor in dados.items():
            if chave.startswith(prefixo):
                celula = linha.cells[1]
                for par in celula.paragraphs[1:]:
                    par._p.getparent().remove(par._p)
                par = celula.paragraphs[0]
                for run in par.runs:
                    run.text = ""
                Montador.escrever(par, valor)


# ------------------------------------------------------------- dados

def _ler_json(caminho: Path) -> Any:
    return json.loads(caminho.read_text(encoding="utf-8")) if caminho.exists() else None


def _ler_csv(caminho: Path) -> list[dict[str, str]]:
    if not caminho.exists():
        return []
    with caminho.open(newline="", encoding="utf-8") as arquivo:
        return list(csv.DictReader(arquivo))


def _d(res: dict[str, Any], rotulo: str) -> dict[str, Any]:
    return res["rq01_04"].get(rotulo) or {"n": 0, "mediana": None, "q1": None, "q3": None}


def _faixa(valor: float | None, casas: int, unidade: str) -> str:
    return f"{fmt(valor, casas)} {unidade}".strip()


def _iqr(d: dict[str, Any], casas: int, unidade: str = "") -> str:
    return (f"mediana de {_faixa(d['mediana'], casas, unidade)} "
            f"(IQR {fmt(d['q1'], casas)}–{fmt(d['q3'], casas)})")


def _nivel(funcao, valor) -> str:
    n = funcao(valor) if valor is not None else None
    return NOMES.get(n, "sem nível") if n else "sem nível"


def _veredito(confirmada: bool | None) -> str:
    return {True: "confirmada", False: "refutada", None: "inconclusiva"}[confirmada]


# ----------------------------------------------------------- conteúdo

def gerar(
    resultados_dir: Path,
    data_dir: Path,
    validacao_dir: Path,
    destino: Path,
    template: Path = TEMPLATE_PADRAO,
    dataset: Path | None = None,
) -> Path:
    res = _ler_json(resultados_dir / "resultados.json")
    if res is None:
        raise FileNotFoundError(
            f"{resultados_dir / 'resultados.json'} não existe: rode `python -m src.pipeline analise`.")
    df = pd.read_csv(dataset or data_dir / "dataset" / "dataset_repos.csv")
    funil = _ler_json(data_dir / "funil.json")
    concord = _ler_json(validacao_dir / "concordancia.json")
    versoes = _ler_csv(validacao_dir / "heuristica_versoes.csv")
    escolhida = _ler_json(validacao_dir / "heuristica_escolhida.json")
    figs = resultados_dir / "figuras"
    n = res["n_repositorios"]

    doc = Document(str(template))
    preencher_capa(doc, {
        "Laboratório": "Lab03 — Mineração de Métricas DORA",
        "Grupo": "[[Integrante 1]] · [[Integrante 2]] · [[Integrante 3]]",
        "Link": "[[link do repositório e do GitHub Projects]]",
        "Data": date.today().strftime("%d/%m/%Y"),
    })
    titulos = limpar_template(doc)
    m = Montador(doc)

    # números usados em várias seções
    df_ = _d(res, "Deployment frequency")
    lta, ltb = _d(res, "Lead time (a) por release"), _d(res, "Lead time (b) por commit")
    cfa, cfb = _d(res, "CFR (a) proxy de CI"), _d(res, "CFR (b) proxy de entrega")
    rec = _d(res, "Tempo de recuperação")
    rq5 = {r["variante"][:7]: r for r in res["rq05"]}
    pares = {p["par"]: p for p in res["rq07"]["pares"]}
    testes = res["rq06"]
    sig = [t for t in testes if t.get("significativo")]
    cats = res["rq01_04"]["categorias_c1"]
    cat_top = max(cats, key=cats.get) if cats else None
    janela = (funil or {}).get("janela", {})
    ini = (janela.get("inicio") or "2025-10-02")[:10]
    fim = (janela.get("fim") or "2026-10-02")[:10]
    data_br = lambda s: f"{s[8:10]}/{s[5:7]}/{s[:4]}"  # noqa: E731
    heuristica_txt = (f"{escolhida['versao']} (F1 = {fmt(escolhida['f1'], 2)})" if escolhida
                      else "[[versão escolhida e F1 — rodar `heuristica`]]")

    # ------------------------------------------------------- 1. Introdução
    m.em(_titulo(titulos, "1."))
    m.p("As métricas DORA (DevOps Research and Assessment) são hoje o padrão de mercado para "
        "medir o desempenho de entrega de software: duas medem velocidade (frequência de deploy e "
        "lead time for changes) e duas medem estabilidade (change failure rate e tempo de "
        "recuperação) [Forsgren et al. 2018; DORA 2024]. O GitHub, porém, não registra deploys "
        "nem falhas em produção. Por isso este laboratório minera proxies públicos — release "
        "publicada como deploy, execução de CI com falha como falha — em repositórios "
        "open-source populares que usam GitHub Actions, e pergunta também o quanto se pode "
        "confiar nesses valores.")
    m.p("Questões de pesquisa do enunciado e hipóteses informais do grupo, formuladas antes da coleta:",
        negrito=True)
    hip = [
        ("RQ1", "Qual a frequência de deploys dos repositórios populares que usam CI/CD?",
         "mediana abaixo de uma release por semana (faixa Medium), pois bibliotecas publicam em ciclos."),
        ("RQ2", "Qual o tempo entre um commit e seu respectivo deploy?",
         "lead time por release (a) bem maior que por commit (b), porque um commit antigo domina a variante (a)."),
        ("RQ3", "Qual a taxa de falha das mudanças entregues?",
         "CFR de CI (a) maior que o CFR de entrega (b): falhas de pipeline são frequentes e corrigidas antes da release."),
        ("RQ4", "Qual o tempo de recuperação após uma execução de CI/CD com falha?",
         "recuperação mediana de poucas horas (menos de 1 dia), pois CI quebrado no branch principal bloqueia o time."),
        ("RQ5", "Repositórios com maior frequência de deploy têm maior ou menor taxa de falha?",
         "correlação fraca ou nula, em linha com o DORA (velocidade e estabilidade não são trade-off)."),
        ("RQ6", "Quais características dos repositórios estão associadas a melhor desempenho DORA?",
         "mais contribuidores, mais releases; diferenças pequenas entre linguagens."),
        ("RQ7", "O quanto a classificação DORA depende da definição operacional?",
         "uma parte relevante dos repositórios (mais de 20%) muda de categoria conforme a definição."),
    ]
    for rq, pergunta, h in hip:
        m.item(f"{rq} — {pergunta} Hipótese: {h}")
    m.p("Além do enunciado (detalhes na seção 3.6), o grupo propõe: a RQ8 (rework rate, métrica "
        "adicionada pelo DORA em 2024) com teste pareado contra o CFR (b); duas combinações extras "
        "na análise de sensibilidade que mudam uma definição por vez; e uma arquitetura de coleta "
        "com cache segmentado por intervalo e coleta paralela em fatias.")

    # --------------------------------------------------------- 2. Contexto
    m.em(_titulo(titulos, "2."))
    m.p("Este é o Lab03 da disciplina, desenvolvido em três sprints e uma entrega final: a S01 "
        "construiu o pipeline de coleta para 100 repositórios; a S02 completou a amostra, as "
        "métricas e a validação manual; a S03, objeto deste relatório, faz a análise estatística "
        "das RQ1–RQ7. Os dados e o pipeline serão reutilizados no Lab04 e replicados por outro "
        "grupo na entrega final.")
    m.p("O objeto de estudo são repositórios públicos do GitHub com mais de 1.000 estrelas que "
        "usam GitHub Actions e publicam releases. A base teórica é o livro Accelerate "
        "[Forsgren et al. 2018] e o guia oficial das quatro métricas [DORA 2024]; a classificação "
        "Elite/High/Medium/Low usa os cortes fixos adaptados pela disciplina. Para a análise "
        "estatística e as ameaças à validade seguimos Wohlin et al. (2012), com as faixas de "
        "Landis e Koch (1977) para kappa e de Romano et al. (2006) para o Cliff's delta.")

    # ---------------------------------------------------- 3.1 Desafios
    m.em(_titulo(titulos, "3.1"))
    for texto in [
        "Rate limit e volume: a API REST permite 5.000 requisições por hora por token, e alguns "
        "repositórios têm centenas de milhares de execuções de CI em 12 meses. A coleta foi dividida "
        "em três fatias, uma por integrante e por token.",
        "Desempenho do cache: na primeira versão, cada página baixada regravava o arquivo inteiro do "
        "repositório; um repositório muito ativo chegou a 185 MB e travou a coleta por horas, embora "
        "só ~100 chamadas por hora estivessem sendo feitas. A solução está na seção 3.6.",
        "Tetos da API: a busca devolve no máximo 1.000 repositórios por consulta, e a listagem de "
        "workflow runs no máximo 1.000 por consulta filtrada; ambos exigiram fatiamento.",
        "Ambiguidade das definições: \"release corretiva\" e \"entrega real\" exigem interpretação "
        "humana, o que motivou a validação com três avaliadores independentes.",
        "Dados incompletos: tags apagadas (compare 404/422), falhas não recuperadas dentro da janela "
        "e releases nos últimos 7 dias exigiram regras explícitas de censura.",
    ]:
        m.item(texto)

    # ---------------------------------------------- 3.2 Tomadas de decisão
    m.em(_titulo(titulos, "3.2"))
    for texto in [
        f"Janela de observação: 12 meses, de {data_br(ini)} a {data_br(fim)}, autorizada pelo professor.",
        "Commits de cada release via compare/{anterior}...{R}, e não por filtro de data: um commit "
        "escrito antes da release anterior, mas só incluído agora, continua contando (é o caso que "
        "faz o lead time (a) crescer). A troca é custo: uma chamada por release.",
        "Ordem de descarte da mais barata para a mais cara (Actions → releases → runs → compare), "
        "para não gastar cota com repositórios que seriam descartados.",
        "Amostra final: os aceitos com mais estrelas, até 400 (folga de ~30% sobre o mínimo de 300 "
        "para absorver erros de coleta).",
        "Classificação DORA com métrica ausente: a métrica fica fora da mediana e exigimos ao menos "
        "3 das 4 métricas; repositórios com menos são reportados como não classificáveis.",
        "Estatística: mediana e IQR em todas as RQs; testes não paramétricos (Spearman, "
        "Kruskal-Wallis, Mann-Whitney) com correção de Holm e tamanho de efeito.",
        "Consenso da amostra-ouro: maioria simples; casos 1-1-1 decididos em reunião e registrados "
        "em desempate.csv.",
        "Processo: board com colunas Backlog → To Do → Doing → Review → Done e limite de WIP de "
        "[[n]] cartões em Doing por integrante, porque [[justificativa]].",
    ]:
        m.item(texto)

    # ------------------------------------------------------- 3.3 Etapas
    m.em(_titulo(titulos, "3.3"))
    m.p("O desenvolvimento seguiu as sprints do enunciado. A divisão abaixo deve refletir os "
        "Assignees reais das Issues no GitHub Projects.")
    m.tabela(
        ["Sprint", "Entregas", "Responsável(is)", "Issues (nº)"],
        [
            ["Lab03S01", "Pipeline de coleta (100 repositórios), cache e rate limit, testes e CI; "
             "Introdução com hipóteses", "[[A, B, C]]", "[[#]]"],
            ["Lab03S02", "Amostra completa (≥ 300), dataset e dicionário de dados, amostra-ouro, "
             "kappa de Fleiss, F1 da heurística; Metodologia", "[[A, B, C]]", "[[#]]"],
            ["Lab03S03", "Classificação DORA, análises RQ1–RQ8, gráficos e relatório",
             "A: RQ1–RQ4 · B: RQ5–RQ6 · C: RQ7–RQ8", "[[#]]"],
            ["Entrega final", "Replicação cruzada, resposta às Issues, ameaças à validade",
             "[[A, B, C]]", "[[#]]"],
        ],
        "Sprints, entregas e responsáveis.", [2.6, 8.4, 3.2, 2.3])
    m.p("Configuração do processo: colunas Backlog → To Do → Doing → Review → Done; limite de WIP "
        "de [[n]] em Doing; todo commit referencia o número da Issue.")
    m.p("[[inserir aqui o print do board do GitHub Projects ao final do laboratório]]")

    # ---------------------------------------------------- 3.4 Ferramentas
    m.em(_titulo(titulos, "3.4"))
    for texto in [
        "Coleta: API REST do GitHub (versão 2026-03-10), acessada por script próprio em Python 3.12 "
        "com requests — sem bibliotecas de terceiros de acesso à API.",
        "Processamento e estatística: pandas, NumPy, SciPy (Spearman, Kruskal-Wallis, Mann-Whitney, "
        "Wilcoxon), statsmodels (kappa de Fleiss, correção de Holm) e scikit-learn (precisão/recall/F1, "
        "kappa de Cohen ponderado).",
        "Visualização: Matplotlib (violino, dispersão, mapa de calor e radar).",
        "Rotulagem manual: tools/rotulador.html, página offline gerada pelo grupo.",
        "Qualidade: pytest e pytest-cov (cobertura mínima de 80% no módulo de métricas), CI no "
        "GitHub Actions a cada push.",
        "Relatório: python-docx, gerando este documento a partir do template da disciplina.",
        "Processo: GitHub Projects (v2) — [[link do board]].",
    ]:
        m.item(texto)

    # ----------------------------------------------- 3.5 Tabela de métricas
    m.em(_titulo(titulos, "3.5"))
    fonte = "Script próprio (API REST)"
    m.tabela(
        ["RQ", "Métrica", "Definição operacional", "Unidade", "Ferramenta / fonte"],
        [
            ["RQ1", "Deployment frequency", "Releases (draft=false, prerelease=false) publicadas na "
             "janela ÷ semanas da janela (≈ 52,2)", "releases/semana", fonte + " — /releases"],
            ["RQ2 (a)", "Lead time por release", "Mediana, entre releases R, de published_at(R) − "
             "commit.author.date do commit mais antigo de compare/{anterior}...{R}", "horas (dias no texto)",
             fonte + " — /compare"],
            ["RQ2 (b)", "Lead time por commit", "Mediana, sobre todos os commits de todas as releases, "
             "de published_at(R) − commit.author.date", "horas (dias no texto)", fonte + " — /compare"],
            ["RQ3 (a)", "CFR — proxy de CI", "Runs com falha (failure, timed_out, startup_failure) ÷ "
             "(falhas + sucessos); push no default branch, na janela", "%", fonte + " — /actions/runs"],
            ["RQ3 (b)", "CFR — proxy de entrega", "Releases seguidas em até 7 dias por release corretiva "
             f"(heurística {heuristica_txt}) ÷ releases avaliadas; últimos 7 dias censurados", "%",
             fonte + " + heurística validada"],
            ["RQ4", "Tempo de recuperação", "Por workflow: updated_at do próximo sucesso − "
             "run_started_at da 1ª falha após um sucesso; mediana dos episódios; sem sucesso = censurado",
             "horas", fonte + " — /actions/runs"],
            ["RQ5", "Correlação velocidade × estabilidade", "ρ de Spearman entre RQ1 e RQ3 (a) e (b)",
             "ρ, p, n", "SciPy"],
            ["RQ6", "Diferença entre grupos", "Kruskal-Wallis (ε² = H/(n−1)) ou Mann-Whitney (Cliff's δ); "
             "Holm sobre todos os testes", "p, ε², δ", "SciPy, statsmodels"],
            ["RQ7", "Sensibilidade", "% de repositórios que mudam de categoria e kappa de Cohen "
             "ponderado (linear) entre as combinações C1–C4", "%, κ", "scikit-learn"],
            ["RQ8*", "Rework rate", "Releases corretivas ÷ releases da janela", "%", "Heurística validada"],
        ],
        "Métricas, definições operacionais e fontes (* = proposta do grupo).",
        [1.6, 2.8, 7.0, 2.2, 3.0])

    # ------------------------------------------------ 3.6 Inovações
    m.em(_titulo(titulos, "3.6"))
    inovacoes = [
        ("(c) Arquitetura de coleta",
         "cache segmentado por intervalo (um arquivo por mês de runs e por compare), coleta "
         "paralela em três fatias (um token por integrante) com junção posterior dos caches, e "
         "agendamento contínuo (um repositório lento não bloqueia os demais). Em teste com 11.700 "
         "runs sintéticos, o tempo de gravação caiu de 8,2 s para 0,7 s; na coleta real, eliminou o "
         "travamento descrito em 3.1. Resultado: seção 4.1."),
        ("(c) Subdivisão adaptativa do teto de 1.000 runs",
         "além de dividir a janela em meses (pedido do enunciado), um mês que atinge o teto é "
         "dividido ao meio, com data e hora, até caber; o que não cabe nem em 1 hora é registrado no "
         "dataset. Resultado: seção 4.1."),
        ("(d) Sensibilidade com uma definição por vez",
         "às combinações do enunciado somamos C3 (só o lead time muda) e C4 (só o CFR muda), "
         "o que permite atribuir a mudança de categoria a cada definição. Resultado: RQ7 em 4.2 e 4.3."),
        ("(a/b) RQ8 — rework rate",
         "métrica adicionada pelo DORA em 2024, comparada ao CFR (b) com Spearman e teste pareado de "
         "Wilcoxon. Resultado: RQ8 em 4.2 e 4.3."),
        ("(d) Heurística versionada e rastreável",
         "cinco versões da heurística de release corretiva avaliadas contra o consenso humano, "
         "sem apagar as anteriores, com o F1 de cada uma reportado. Resultado: tabela de validação em 4.1."),
    ]
    for titulo, texto in inovacoes:
        m.item(f"{titulo}: {texto}")

    # ------------------------------------------------ 4.1 Coleta de dados
    m.em(_titulo(titulos, "4.1"))
    if funil:
        m.p(f"A busca reuniu {fmt(funil['etapas'][0]['n'])} candidatos com mais de 1.000 estrelas. "
            f"A Tabela {m.n_tabela + 1} mostra quantos restaram em cada etapa. A amostra final tem "
            f"{fmt(funil['amostra_final'])} repositórios "
            f"({'atinge' if funil['atinge_minimo'] else 'NÃO atinge'} o mínimo de "
            f"{funil['minimo_exigido']}), com janela de {data_br(ini)} a {data_br(fim)}.")
        m.tabela(["Etapa", "Repositórios", "Descartados", "Motivo"],
                 [[e["etapa"], fmt(e["n"]), fmt(e["descartados"]), e["motivo"] or "—"]
                  for e in funil["etapas"]],
                 "Funil de seleção dos repositórios.", [4.6, 2.4, 2.4, 7.1])
    else:
        m.p(f"[[funil de seleção — rodar `consolidate` e gerar data/funil.json]] A amostra analisada "
            f"tem {fmt(n)} repositórios.")
    ep, cens = int(df["rec_n_episodios"].sum()), int(df["rec_n_censurados"].sum())
    aval, rcens = int(df["cfr_n_avaliadas"].sum()), int(df["cfr_n_censuradas"].sum())
    trunc = int((df.get("runs_intervalos_truncados", pd.Series([0])) > 0).sum())
    m.p(f"Ao todo foram analisadas {fmt(int(df['n_releases'].sum()))} releases estáveis e "
        f"{fmt(int((df['ci_n_sucessos'] + df['ci_n_falhas']).sum()) if df['ci_n_sucessos'].notna().any() else 0)} "
        f"execuções de CI válidas. Dados ausentes e censurados não foram descartados: "
        f"{fmt(cens)} de {fmt(ep)} episódios de falha ({fmt(100 * cens / ep if ep else None, 1)}%) "
        f"terminaram sem recuperação na janela (censurados); {fmt(rcens)} releases dos últimos 7 dias "
        f"ficaram fora do CFR (b); {fmt(int(df['lt_n_releases_compare_404'].sum()))} releases foram "
        f"ignoradas no lead time por compare 404/422; {fmt(trunc)} repositórios tiveram algum intervalo "
        f"de runs truncado pelo teto da API. Outliers foram mantidos: todas as estatísticas usam "
        f"mediana e IQR, robustas a valores extremos.")
    # validação manual
    m.p("Validação manual (amostra-ouro): 60 repositórios sorteados (random_state = 42) e até 5 "
        "releases de cada, rotulados de forma independente pelos três integrantes, sem ver a "
        "heurística.", negrito=False)
    if concord:
        nomes_dim = {"tipo_projeto": "Tipo do projeto", "entregas_reais": "Entregas reais",
                     "corretiva": "Release corretiva"}
        m.tabela(["Dimensão", "Itens", "Concordância total", "Kappa de Fleiss", "Interpretação"],
                 [[nomes_dim.get(k, k), fmt(v["n_itens"]),
                   f"{fmt(100 * v['concordancia_total'], 1)}%" if v.get("concordancia_total") is not None else "—",
                   fmt(v.get("kappa_fleiss"), 3), v.get("interpretacao") or "—"]
                  for k, v in concord["dimensoes"].items()],
                 "Concordância entre os três avaliadores (Landis e Koch, 1977)."
                 + (" Rotulagem parcial." if concord.get("parcial") else ""),
                 [4.0, 1.8, 3.4, 3.2, 4.1])
        m.p(f"Casos sem maioria decididos em reunião: {concord.get('casos_sem_maioria', 0)}.")
    else:
        m.p("[[tabela de kappa de Fleiss — rodar `python -m src.pipeline concordancia`]]")
    if versoes:
        m.tabela(["Versão", "Precisão", "Recall", "F1", "n"],
                 [[v["versao"], fmt(float(v["precisao"]), 3), fmt(float(v["recall"]), 3),
                   fmt(float(v["f1"]), 3), v["n"]] for v in versoes],
                 f"Versões da heurística de release corretiva frente ao consenso; escolhida: {heuristica_txt}.",
                 [5.0, 2.8, 2.8, 2.8, 2.0])
        if escolhida and not escolhida.get("atinge_meta_f1_0_70"):
            m.p("O melhor F1 ficou abaixo da meta de 0,70: [[descrever a versão refinada (v6) e o novo F1]].")
    else:
        m.p("[[tabela de precisão/recall/F1 — rodar `python -m src.pipeline heuristica`]]")

    # ------------------------------------------- 4.2 Visualização gráfica
    m.em(_titulo(titulos, "4.2"))
    m.p("Todas as medidas abaixo são calculadas por repositório e resumidas pela mediana e pelo "
        "intervalo interquartil (IQR), porque as distribuições são muito assimétricas.")
    m.h3("RQ1–RQ4: qual o desempenho DORA dos repositórios?")
    m.figura(figs / "rq01_04_distribuicoes.png",
             "Distribuição das quatro métricas DORA por repositório (violino; ponto = mediana, barra = IQR).",
             largura_cm=15)
    m.p(f"RQ1: {_iqr(df_, 2, 'releases/semana')}, faixa {_nivel(nivel_deploy_freq, df_['mediana'])}. "
        f"RQ2: lead time (a) com {_iqr(lta, 1, 'dias')} e (b) com {_iqr(ltb, 1, 'dias')}. "
        f"RQ3: CFR de CI (a) com {_iqr(cfa, 1, '%')} e CFR de entrega (b) com {_iqr(cfb, 1, '%')}. "
        f"RQ4: recuperação com {_iqr(rec, 1, 'horas')}.")
    tab14 = _ler_csv(resultados_dir / "tab_rq01_04.csv")
    if tab14:
        m.tabela(list(tab14[0].keys()), [list(r.values()) for r in tab14],
                 "Mediana e IQR das métricas por repositório.", [1.6, 4.4, 2.6, 1.3, 1.9, 2.2, 1.9])
    tabdora = _ler_csv(resultados_dir / "tab_dora_c1.csv")
    if tabdora:
        m.tabela(list(tabdora[0].keys()), [list(r.values()) for r in tabdora],
                 "Classificação DORA na combinação de referência C1 (release, lead time (a), CFR (a)).",
                 [3.6, 2.4, 2.4, 2.4, 2.4, 2.2])
    m.h3("RQ5: repositórios que entregam mais falham mais?")
    m.figura(figs / "rq05_dispersao.png",
             "Frequência de deploy (escala log) × CFR, por repositório, com o ρ de Spearman.")
    a, b = rq5.get("CFR (a)", {}), rq5.get("CFR (b)", {})
    m.p(f"CFR de CI: ρ = {fmt(a.get('rho'), 3)} (p = {fmt_p(a.get('p'))}, n = {a.get('n')}); "
        f"CFR de entrega: ρ = {fmt(b.get('rho'), 3)} (p = {fmt_p(b.get('p'))}, n = {b.get('n')}).")
    m.h3("RQ6: quais características estão associadas a melhor desempenho?")
    m.figura(figs / "rq06_efeitos.png",
             "Tamanho de efeito de cada fator sobre cada métrica (* = significativo após Holm).",
             largura_cm=15)
    m.p(f"Foram feitos {len(testes)} testes; {len(sig)} permaneceram significativos após a correção "
        f"de Holm" + (": " + "; ".join(
            f"{t['fator']} × {t['metrica']} ({t['efeito']}, efeito {t['magnitude']})" for t in sig)
            if sig else ".") + " As medianas por grupo estão em rq06_medianas_por_grupo.csv.")
    m.h3("RQ7: a classificação muda com a definição?")
    m.figura(figs / "rq07_radar.png",
             "% de repositórios Elite ou High por métrica: cada combinação (azul) sobre a referência C1 (cinza).")
    linhas7 = [[p["par"], fmt(p["n"]),
                f"{fmt(100 * p['mudaram'], 1)}%" if p.get("mudaram") is not None else "—",
                fmt(p.get("kappa"), 3)] for p in res["rq07"]["pares"]]
    m.tabela(["Par", "n", "Mudaram de categoria", "Kappa ponderado"], linhas7,
             "Concordância entre as combinações C1 (release, (a), (a)), C2 (release + pré-release, (b), (b)), "
             "C3 (release, (b), (a)) e C4 (release, (a), (b)).", [3.4, 2.0, 4.6, 4.0])
    r8 = res["rq08"]
    m.h3("RQ8 (proposta do grupo): rework rate × CFR (b)")
    tab8 = _ler_csv(resultados_dir / "tab_rq08.csv")
    m.p((f"Rework rate: {tab8[0]['Mediana']}% de mediana (IQR {tab8[0]['Q1']}–{tab8[0]['Q3']}%). " if tab8 else "")
        + f"Spearman com o CFR (b): ρ = {fmt(r8.get('rho'), 3)} (p = {fmt_p(r8.get('p_rho'))}, n = {r8.get('n')}); "
        f"Wilcoxon pareado: p = {fmt_p(r8.get('p_wilcoxon'))}.")

    # ------------------------------------------------------- 4.3 Discussão
    m.em(_titulo(titulos, "4.3"))
    h = {
        1: df_["mediana"] is not None and df_["mediana"] < 1,
        2: None if lta["mediana"] is None or ltb["mediana"] is None else lta["mediana"] > ltb["mediana"],
        3: None if cfa["mediana"] is None or cfb["mediana"] is None else cfa["mediana"] > cfb["mediana"],
        4: rec["mediana"] is not None and rec["mediana"] < 24,
        5: all(r.get("rho") is None or abs(r["rho"]) < 0.3 or r["p"] >= 0.05 for r in res["rq05"]),
        7: (pares.get("C1 × C2") or {}).get("mudaram") is not None and pares["C1 × C2"]["mudaram"] > 0.2,
    }
    contrib_df = next((t for t in testes if t["fator"].startswith("Contribuidores")
                       and t["metrica"] == "Deployment frequency"), None)
    h[6] = None if contrib_df is None else bool(contrib_df["significativo"])
    m.p("Hipóteses × resultados:", negrito=True)
    m.item(f"RQ1 — {_veredito(h[1])}: a mediana foi de {fmt(df_['mediana'], 2)} releases/semana "
           f"(faixa {_nivel(nivel_deploy_freq, df_['mediana'])}).")
    m.item(f"RQ2 — {_veredito(h[2])}: (a) {fmt(lta['mediana'], 1)} dias × (b) {fmt(ltb['mediana'], 1)} dias. "
           "A variante (a) é determinada pelo commit mais antigo da release: basta um commit escrito meses antes "
           "e incluído só agora (um ramo longo, um backport, um PR esquecido) para a release inteira parecer lenta; "
           "a variante (b) dilui esse commit entre todos os demais.")
    m.item(f"RQ3 — {_veredito(h[3])}: CFR (a) {fmt(cfa['mediana'], 1)}% × CFR (b) {fmt(cfb['mediana'], 1)}%. "
           "O CFR (a) mede falha de pipeline — testes instáveis, quebras de dependência e de infraestrutura entram "
           "como falha sem que nenhum usuário as veja —, enquanto o CFR (b) se aproxima da definição do DORA "
           f"(um deploy que exigiu correção), mas depende da heurística ({heuristica_txt}).")
    m.item(f"RQ4 — {_veredito(h[4])}: recuperação mediana de {fmt(rec['mediana'], 1)} horas "
           f"(faixa {_nivel(nivel_recuperacao, rec['mediana'])}); "
           f"{fmt(100 * cens / ep if ep else None, 1)}% dos episódios foram censurados, e descartá-los faria os "
           "repositórios parecerem mais rápidos do que são.")
    m.item(f"RQ5 — {_veredito(h[5])}: ρ = {fmt(a.get('rho'), 2)} (CI) e {fmt(b.get('rho'), 2)} (entrega). "
           + ("Não há evidência de trade-off entre velocidade e estabilidade, o que é coerente com o DORA. "
              if h[5] else
              "Há associação entre entregar mais e falhar mais neste proxy, o que contradiz a afirmação do DORA. ")
           + "Correlação não implica causalidade: projetos maiores podem ter, ao mesmo tempo, mais releases e mais CI.")
    m.item(f"RQ6 — {_veredito(h[6])}: "
           + (f"contribuidores × frequência de deploy teve p ajustado = {fmt_p(contrib_df['p_holm'])} "
              f"({contrib_df['efeito']}, efeito {contrib_df['magnitude']}). " if contrib_df else "")
           + (f"{len(sig)} de {len(testes)} testes foram significativos após Holm." if testes else ""))
    c12, c13, c14 = (pares.get(k) or {} for k in ("C1 × C2", "C1 × C3", "C1 × C4"))
    pct = lambda p: f"{fmt(100 * p['mudaram'], 0)}%" if p.get("mudaram") is not None else "—"  # noqa: E731
    m.item(f"RQ7 — {_veredito(h[7])}: {pct(c12)} dos repositórios mudam de categoria entre C1 e C2 "
           f"(κ = {fmt(c12.get('kappa'), 2)}). Mudando só o lead time (C1 × C3), {pct(c13)} mudam; "
           f"mudando só o CFR (C1 × C4), {pct(c14)}. "
           + ("A classificação DORA de um repositório open-source depende fortemente da definição escolhida e "
              "deve ser lida com cautela." if h[7] else
              "A classificação é relativamente robusta à definição."))
    m.p("O que surpreendeu: [[1–2 achados inesperados e uma explicação plausível]].")
    m.p("Ameaças à validade:", negrito=True)
    kappa_corr = ((concord or {}).get("dimensoes", {}).get("corretiva") or {}).get("kappa_fleiss")
    for texto in [
        "Construto: release não é deploy (uma biblioteca publica versões que só chegam a produção quando os "
        "usuários atualizam) e falha de CI não é falha em produção; a data de autor do commit é distorcida por "
        f"rebase e squash merge; a dimensão \"corretiva\" teve kappa de Fleiss = {fmt(kappa_corr, 2) if kappa_corr is not None else '[[kappa]]'}.",
        "Interna: repositórios mais antigos e com mais contribuidores tendem a ter mais releases e mais CI; o "
        "compare entre releases de ramos de manutenção pode incluir commits de outro ramo; cada fatia foi "
        "coletada em um momento diferente.",
        "Externa: só repositórios populares (> 1.000 estrelas), públicos e que usam releases e GitHub Actions; "
        "não vale para projetos corporativos ou que publicam por tags e registros de pacotes. A variante \"tag\" "
        "não foi coletada.",
        f"Conclusão: {len(testes)} testes na RQ6 (controlados por Holm); grupos pequenos em alguns fatores; a "
        "melhor versão da heurística foi escolhida no mesmo conjunto em que foi avaliada (F1 otimista).",
    ]:
        m.item(texto)
    m.p("Contribuição das inovações: a sensibilidade com uma definição por vez (C3 e C4) mostrou qual definição "
        "mais altera a classificação; a RQ8 mostra que rework rate e CFR (b), embora usem a mesma heurística, "
        f"medem coisas diferentes (ρ = {fmt(r8.get('rho'), 2)}); e a arquitetura de coleta tornou viável a "
        "amostra completa dentro do prazo.")

    # ------------------------------------------------------- 5. Conclusão
    m.em(_titulo(titulos, "5. Conclus"))
    m.p(f"Nos {fmt(n)} repositórios analisados, a categoria DORA mais comum na definição de referência foi "
        f"{cat_top or '—'}. Velocidade e estabilidade "
        + ("não mostraram trade-off" if h[5] else "mostraram associação")
        + ", e a classificação mostrou-se "
        + ("sensível" if h[7] else "pouco sensível")
        + " à definição operacional — a principal mensagem é que métricas DORA mineradas do GitHub são proxies "
        "e precisam ser reportadas junto com a definição usada.")
    m.p("Limitações: proxies de deploy e falha, heurística de release corretiva com F1 de "
        f"{fmt(escolhida['f1'], 2) if escolhida else '[[F1]]'}, amostra restrita a repositórios populares e "
        "variante \"tag\" não coletada.")
    m.p("Com mais tempo, o grupo coletaria a variante por tags (datas via GraphQL) e o compare das "
        "pré-releases, separaria um conjunto de teste para a heurística e expandiria a RQ8 para a evolução "
        "trimestral das métricas (Mann-Kendall).")

    # --------------------------------------------------- Referências
    ref = _titulo(titulos, "5. Refer") if any(t.startswith("5. Refer") for t in titulos) else _titulo(titulos, "6.")
    for el_texto in ref.iter(qn("w:t")):
        el_texto.text = el_texto.text.replace("5. Referências", "6. Referências")
    m.em(ref)
    for texto in [
        "DORA. DORA's software delivery metrics: the four keys. 2024. Disponível em: https://dora.dev/guides/dora-metrics-four-keys/.",
        "DORA. A history of DORA's software delivery metrics. Disponível em: https://dora.dev/insights/dora-metrics-history/.",
        "FORSGREN, N.; HUMBLE, J.; KIM, G. Accelerate: The Science of Lean Software and DevOps. IT Revolution, 2018.",
        "HOLM, S. A simple sequentially rejective multiple test procedure. Scandinavian Journal of Statistics, v. 6, n. 2, p. 65–70, 1979.",
        "LANDIS, J. R.; KOCH, G. G. The measurement of observer agreement for categorical data. Biometrics, v. 33, n. 1, p. 159–174, 1977.",
        "ROMANO, J. et al. Appropriate statistics for ordinal level data. Annual Meeting of the Florida Association of Institutional Research, 2006.",
        "WOHLIN, C. et al. Experimentation in Software Engineering. Springer, 2012.",
        "ZUSE, H. A framework of software measurement. Walter de Gruyter, 2013.",
    ]:
        m.p(texto, "List Paragraph")

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return destino
