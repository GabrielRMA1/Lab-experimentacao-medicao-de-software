"""Relatório .docx gerado a partir do modelo da disciplina."""

from docx import Document

from src.analise import executar_tudo
from src.analise.relatorio import gerar
from tests.dados_sinteticos import preparar_ambiente


def test_relatorio_docx_no_modelo_da_disciplina(tmp_path):
    csv, saida, validacao = preparar_ambiente(tmp_path)
    executar_tudo(csv, saida, validacao)
    destino = gerar(saida, csv.parent, validacao, saida / "Relatorio.docx", dataset=csv)
    doc = Document(str(destino))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "ORIENTAÇÃO" not in texto and "[conteúdo do grupo" not in texto
    assert "4.2 Visualização Gráfica" in texto and "6. Referências" in texto
    assert "RQ1 —" in texto and "[PREENCHER:" in texto  # marcadores para o grupo
    assert len(doc.inline_shapes) == 4  # uma figura por bloco de RQs
    assert doc.tables[0].rows[4].cells[1].text.startswith("Lab03")
